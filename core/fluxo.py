"""Fluxo de caixa do mês a partir da mesma base classificada da DRE.

Estrutura: SALDO INICIAL → seções da DRE (grupos e contas) → transferências entre contas →
VARIAÇÃO DO CAIXA → SALDO FINAL. Cada linha tem o valor realizado e, opcionalmente, os
"gastos previstos" do mês.

Previsto de cada conta: o valor digitado pelo usuário para o mês ou, se não houver, a sugestão
realizado + lançamentos ainda em aberto no OMIE no mês (previstos, atrasados, a vencer).
"""
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from . import dre as D
from .excel import AZUL, AZUL_CLARO, BORDA, CINZA, CINZA_CLARO, FMT, FMT_PCT, _cab
from .plano import IGNORAR, SECOES

TRANSF = ("TRANSFERENCIAS", "(+/-) TRANSFERÊNCIAS ENTRE CONTAS")


def chave(grupo, conta):
    return f"{grupo}|{conta}"


def base_pendente(lancs, cfg, mapa):
    """Lançamentos do mês ainda não realizados (previstos/atrasados/a vencer), classificados como na DRE."""
    pend = [l for l in lancs if not D.realizado(l)]
    return D.montar_base(pend, dict(cfg, apenas_realizado=False, rateios_mes=[]), mapa) if pend else []


def montar_fluxo(base, pendentes, saldos, cfg, previstos=None):
    """Linhas do fluxo: {nivel, tipo, descricao, chave, realizado, previsto, sugestao}.

    previstos: {chave(grupo, conta): valor} digitados pelo usuário para o mês.
    """
    previstos = previstos or {}
    real, pend = defaultdict(float), defaultdict(float)
    for b in base:
        real[(b["grupo"], b["conta"])] += b["valor"]
    for b in pendentes:
        pend[(b["grupo"], b["conta"])] += b["valor"]
    contas = set(real) | set(pend)

    ini = sum(s.get("saldo_inicial") or 0 for s in saldos or [])
    linhas = [{"nivel": 1, "tipo": "saldo", "descricao": "SALDO INICIAL", "chave": "", "realizado": ini,
               "previsto": ini, "sugestao": ini}]
    grupos = list(D.grupos_da_empresa(cfg))
    secoes = list(SECOES) + [TRANSF]
    tot_real = tot_prev = 0.0
    for sec, rotulo in secoes:
        lin_sec = {"nivel": 1, "tipo": "secao", "descricao": rotulo, "chave": "", "realizado": 0.0, "previsto": 0.0}
        filhos = []
        for cod, nome, s, _rel in grupos:
            if (sec == "TRANSFERENCIAS") != (cod == IGNORAR) or (sec != "TRANSFERENCIAS" and s != sec):
                continue
            cs = sorted((c for g, c in contas if g == cod), key=str.lower)
            if not cs:
                continue
            lin_g = {"nivel": 2, "tipo": "grupo", "descricao": nome if cod == IGNORAR else f"{cod}-{nome}",
                     "chave": "", "realizado": 0.0, "previsto": 0.0}
            filhos.append(lin_g)
            for c in cs:
                k = chave(cod, c)
                r = round(real.get((cod, c), 0.0), 2)
                sug = round(r + pend.get((cod, c), 0.0), 2)
                p = round(float(previstos[k]), 2) if k in previstos else sug
                filhos.append({"nivel": 3, "tipo": "conta", "descricao": c, "chave": k, "realizado": r,
                               "previsto": p, "sugestao": sug})
                lin_g["realizado"] += r
                lin_g["previsto"] += p
            lin_sec["realizado"] += lin_g["realizado"]
            lin_sec["previsto"] += lin_g["previsto"]
        if filhos:
            linhas.append(lin_sec)
            linhas += filhos
            tot_real += lin_sec["realizado"]
            tot_prev += lin_sec["previsto"]
    linhas.append({"nivel": 1, "tipo": "resultado", "descricao": "(=) VARIAÇÃO DO CAIXA", "chave": "",
                   "realizado": tot_real, "previsto": tot_prev})
    linhas.append({"nivel": 1, "tipo": "saldo", "descricao": "SALDO FINAL", "chave": "",
                   "realizado": ini + tot_real, "previsto": ini + tot_prev})
    return linhas


def resumo(linhas):
    """Totais para os indicadores da tela: saldo inicial, entradas, saídas, saldo final (realizado)."""
    contas = [l for l in linhas if l["tipo"] == "conta" and not l["chave"].startswith(IGNORAR + "|")]
    return {"saldo_inicial": linhas[0]["realizado"],
            "entradas": sum(l["realizado"] for l in contas if l["realizado"] > 0),
            "saidas": sum(l["realizado"] for l in contas if l["realizado"] < 0),
            "saldo_final": linhas[-1]["realizado"],
            "saldo_final_previsto": linhas[-1]["previsto"]}


def diario(base, saldos):
    """Movimento por dia (realizado): entradas, saídas, transferências e saldo acumulado."""
    dias = defaultdict(lambda: {"entradas": 0.0, "saidas": 0.0, "transferencias": 0.0})
    for b in base:
        d = dias[b["data"]]
        if b["grupo"] == IGNORAR:
            d["transferencias"] += b["valor"]
        elif b["valor"] >= 0:
            d["entradas"] += b["valor"]
        else:
            d["saidas"] += b["valor"]
    saldo = sum(s.get("saldo_inicial") or 0 for s in saldos or [])
    out = []
    for data in sorted(dias):
        d = dias[data]
        saldo += d["entradas"] + d["saidas"] + d["transferencias"]
        out.append({"data": data, **{k: round(v, 2) for k, v in d.items()}, "saldo": round(saldo, 2)})
    return out


# ---------------------------------------------------------------- Excel
def gerar_excel_fluxo(caminho, empresa, ano, mes, linhas, dias, base, com_previsto=True):
    wb = Workbook()
    ws = wb.active
    ws.title = f"FLUXO DE CAIXA {mes:02d} {ano}"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 52
    ws.cell(1, 2, f"{empresa['nome']} — Fluxo de Caixa {mes:02d}/{ano}").font = Font(bold=True, size=14, color=AZUL)
    cab = ["Descrição", "Realizado"] + (["Gastos previstos", "Diferença", "% Realizado"] if com_previsto else [])
    for i, h in enumerate(cab, start=2):
        _cab(ws, 3, i, h)
        if i > 2:
            ws.column_dimensions[chr(64 + i)].width = 17
    ws.freeze_panes = "C4"
    estilo = {"saldo": (True, AZUL_CLARO), "secao": (True, CINZA), "resultado": (True, CINZA_CLARO),
              "grupo": (True, None), "conta": (False, None)}
    for r, l in enumerate(linhas, start=4):
        negrito, fundo = estilo[l["tipo"]]
        recuo = {"grupo": "   ", "conta": "      "}.get(l["tipo"], "")
        vals = [recuo + l["descricao"], round(l["realizado"], 2)]
        if com_previsto:
            vals += [round(l["previsto"], 2), f"=C{r}-D{r}", f'=IF(D{r}=0,"",C{r}/D{r})']
        for c, v in enumerate(vals, start=2):
            cel = ws.cell(r, c, v)
            cel.border = BORDA
            cel.font = Font(bold=negrito)
            if fundo:
                cel.fill = PatternFill("solid", fgColor=fundo)
            if c > 2:
                cel.number_format = FMT_PCT if c == 6 else FMT
            if c == 4 and l["tipo"] == "conta":
                cel.fill = PatternFill("solid", fgColor="FFF2CC")  # célula editável
    if com_previsto:
        ws.cell(len(linhas) + 5, 2, "Células em amarelo: gastos previstos do mês (podem ser ajustados). "
                                    "Os totais desta coluna foram calculados pelo sistema.").font = Font(italic=True, size=9)

    wd = wb.create_sheet("FLUXO DIÁRIO")
    wd.sheet_view.showGridLines = False
    for i, (h, w) in enumerate(zip(["Data", "Entradas", "Saídas", "Transferências", "Saldo do dia"],
                                   (12, 16, 16, 16, 16)), start=2):
        _cab(wd, 2, i, h)
        wd.column_dimensions[chr(64 + i)].width = w
    for r, d in enumerate(dias, start=3):
        for c, v in enumerate([d["data"], d["entradas"], d["saidas"], d["transferencias"], d["saldo"]], start=2):
            cel = wd.cell(r, c, v)
            cel.border = BORDA
            cel.number_format = "DD/MM/YYYY" if c == 2 else FMT
    wd.freeze_panes = "B3"

    from .excel import _aba_base
    _aba_base(wb, base)
    wb.save(caminho)
    return caminho
