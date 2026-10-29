"""Modelos de DRE: prontos por segmento, importados de um Excel de exemplo e exportados para edição.

Um modelo tem:
  grupos = [[codigo, nome, secao, relatorio], ...]   (nível 2 da DRE, como no plano padrão)
  contas = [[nome da conta, codigo do grupo], ...]   (nível 3; usado para classificar as categorias do OMIE)
"""
import re

import openpyxl
from openpyxl.styles import Font, PatternFill

from .plano import GRUPOS_PADRAO, IGNORAR, SECOES, norm

A_CLASSIFICAR = ["060900", "OUTRAS DESPESAS (A CLASSIFICAR)", "DESPESAS", ""]
TRANSFERENCIAS = list(GRUPOS_PADRAO[-1])

# ---------------------------------------------------------------- modelos prontos
_SERVICOS_CONTAS = [
    ("Serviços Prestados", "010100"), ("Clientes - Serviços Prestados", "010100"),
    ("Impostos s/ Faturamento", "020100"), ("ISS", "020100"), ("Simples Nacional", "020100"),
    ("Salários", "040100"), ("FT Freelancer", "040100"), ("Rescisões", "040100"), ("Adiantamento Salarial", "040100"),
    ("FGTS", "040200"), ("INSS", "040200"), ("Vale Refeição", "040300"), ("Vale Alimentação", "040300"),
    ("Vale Transporte", "040300"), ("Seguro de Vida", "040300"), ("Assistência Médica", "040300"),
    ("Limpeza", "040400"), ("Uniformes", "040400"), ("Combustível", "040500"), ("Locação de Veículos", "040500"),
    ("Compra de Serviços", "040700"), ("Contabilidade", "060100"), ("Sistemas", "060200"),
    ("Telefonia", "060400"), ("Marketing", "060400"), ("Aluguel", "060500"), ("Energia Elétrica", "060500"),
    ("Advogados", "060600"), ("Pró-Labore", "060700"), ("Rendimentos de Aplicações", "080100"),
    ("Tarifas Bancárias", "080200"), ("Aporte Sócio", "090100"), ("Empréstimos", "090200"),
]

_COMERCIO_GRUPOS = [
    ("010100", "VENDAS DE MERCADORIAS", "ENTRADAS", "recebimentos"),
    ("010200", "OUTRAS ENTRADAS OPERACIONAIS", "ENTRADAS", "recebimentos"),
    ("020100", "IMPOSTOS SOBRE VENDAS", "DEDUCOES", ""),
    ("020200", "DEVOLUÇÕES E ESTORNOS", "DEDUCOES", ""),
    ("040100", "COMPRA DE MERCADORIAS (CMV)", "CUSTOS", "fornecedores"),
    ("040200", "FRETES E LOGÍSTICA", "CUSTOS", "fornecedores"),
    ("040300", "COMISSÕES DE VENDAS", "CUSTOS", ""),
    ("040400", "EMBALAGENS", "CUSTOS", "fornecedores"),
    ("050100", "PESSOAL", "DESPESAS", "funcionarios"),
    ("050200", "ENCARGOS E BENEFÍCIOS", "DESPESAS", ""),
    ("060100", "CONTABILIDADE", "DESPESAS", ""),
    ("060200", "SISTEMAS E TECNOLOGIA", "DESPESAS", ""),
    ("060300", "MARKETING E VENDAS", "DESPESAS", "fornecedores"),
    ("060400", "ADMINISTRATIVO", "DESPESAS", ""),
    ("060500", "OCUPAÇÃO (ALUGUEL, CONDOMÍNIO, UTILIDADES)", "DESPESAS", ""),
    ("060700", "SÓCIOS", "DESPESAS", ""),
    tuple(A_CLASSIFICAR),
    ("080100", "RECEITAS FINANCEIRAS", "FINANCEIRO", ""),
    ("080200", "DESPESAS FINANCEIRAS E TAXAS DE CARTÃO", "FINANCEIRO", ""),
    ("090100", "ENTRADAS NÃO OPERACIONAIS", "NAO_OPERACIONAL", ""),
    ("090200", "SAÍDAS NÃO OPERACIONAIS", "NAO_OPERACIONAL", ""),
]
_COMERCIO_CONTAS = [
    ("Vendas", "010100"), ("Venda de Mercadorias", "010100"), ("Vendas Cartão", "010100"), ("Vendas PIX", "010100"),
    ("Simples Nacional", "020100"), ("ICMS", "020100"), ("PIS", "020100"), ("COFINS", "020100"),
    ("Devoluções", "020200"), ("Compra de Mercadorias", "040100"), ("Fornecedores", "040100"),
    ("Frete", "040200"), ("Fretes s/ Compras", "040200"), ("Comissões", "040300"), ("Embalagens", "040400"),
    ("Salários", "050100"), ("Rescisões", "050100"), ("FGTS", "050200"), ("INSS", "050200"),
    ("Vale Transporte", "050200"), ("Vale Refeição", "050200"), ("Contabilidade", "060100"),
    ("Sistemas", "060200"), ("Marketing", "060300"), ("Publicidade", "060300"), ("Telefonia", "060400"),
    ("Aluguel", "060500"), ("Energia Elétrica", "060500"), ("Pró-Labore", "060700"),
    ("Taxa de Cartão", "080200"), ("Antecipação de Recebíveis", "080200"), ("Tarifas Bancárias", "080200"),
]

_EVENTOS_GRUPOS = [
    ("010100", "RECEBIMENTOS DE EVENTOS", "ENTRADAS", "recebimentos"),
    ("010200", "OUTRAS ENTRADAS OPERACIONAIS", "ENTRADAS", "recebimentos"),
    ("020100", "IMPOSTOS", "DEDUCOES", ""),
    ("040100", "EQUIPE DE EVENTOS", "CUSTOS", "funcionarios"),
    ("040200", "ENCARGOS E BENEFÍCIOS", "CUSTOS", ""),
    ("040300", "ALIMENTOS E BEBIDAS", "CUSTOS", "fornecedores"),
    ("040400", "DECORAÇÃO E MATERIAIS", "CUSTOS", "fornecedores"),
    ("040500", "LOCAÇÃO DE ESPAÇO E EQUIPAMENTOS", "CUSTOS", "fornecedores"),
    ("040600", "ATRAÇÕES (DJ, BANDA, SOM E LUZ)", "CUSTOS", "fornecedores"),
    ("040700", "TRANSPORTE E LOGÍSTICA", "CUSTOS", "fornecedores"),
    ("060100", "CONTABILIDADE", "DESPESAS", ""),
    ("060200", "SISTEMAS E TECNOLOGIA", "DESPESAS", ""),
    ("060300", "MARKETING", "DESPESAS", ""),
    ("060400", "ADMINISTRATIVO", "DESPESAS", ""),
    ("060500", "OCUPAÇÃO (ALUGUEL, CONDOMÍNIO, UTILIDADES)", "DESPESAS", ""),
    ("060700", "SÓCIOS", "DESPESAS", ""),
    tuple(A_CLASSIFICAR),
    ("080100", "RECEITAS FINANCEIRAS", "FINANCEIRO", ""),
    ("080200", "DESPESAS FINANCEIRAS", "FINANCEIRO", ""),
    ("090100", "ENTRADAS NÃO OPERACIONAIS", "NAO_OPERACIONAL", ""),
    ("090200", "SAÍDAS NÃO OPERACIONAIS", "NAO_OPERACIONAL", ""),
]
_EVENTOS_CONTAS = [
    ("Eventos", "010100"), ("Festas", "010100"), ("Sinal de Evento", "010100"), ("Clientes", "010100"),
    ("Simples Nacional", "020100"), ("ISS", "020100"), ("Garçons", "040100"), ("Freelancers", "040100"),
    ("Cachê", "040100"), ("Salários", "040100"), ("FGTS", "040200"), ("INSS", "040200"),
    ("Buffet", "040300"), ("Bebidas", "040300"), ("Alimentos", "040300"), ("Decoração", "040400"),
    ("Descartáveis", "040400"), ("Locação de Espaço", "040500"), ("Locação de Equipamentos", "040500"),
    ("DJ", "040600"), ("Banda", "040600"), ("Som e Iluminação", "040600"), ("Frete", "040700"),
    ("Combustível", "040700"), ("Contabilidade", "060100"), ("Sistemas", "060200"), ("Marketing", "060300"),
    ("Aluguel", "060500"), ("Energia Elétrica", "060500"), ("Pró-Labore", "060700"),
    ("Tarifas Bancárias", "080200"), ("Taxa de Cartão", "080200"),
]

_INDUSTRIA_GRUPOS = [
    ("010100", "VENDAS DE PRODUTOS", "ENTRADAS", "recebimentos"),
    ("010200", "OUTRAS ENTRADAS OPERACIONAIS", "ENTRADAS", "recebimentos"),
    ("020100", "IMPOSTOS SOBRE VENDAS (ICMS, IPI, PIS, COFINS)", "DEDUCOES", ""),
    ("020200", "DEVOLUÇÕES", "DEDUCOES", ""),
    ("040100", "MATÉRIA-PRIMA E INSUMOS", "CUSTOS", "fornecedores"),
    ("040200", "MÃO DE OBRA DA PRODUÇÃO", "CUSTOS", "funcionarios"),
    ("040300", "ENCARGOS E BENEFÍCIOS DA PRODUÇÃO", "CUSTOS", ""),
    ("040400", "ENERGIA E UTILIDADES DA FÁBRICA", "CUSTOS", ""),
    ("040500", "MANUTENÇÃO DE MÁQUINAS", "CUSTOS", "fornecedores"),
    ("040600", "FRETES E LOGÍSTICA", "CUSTOS", "fornecedores"),
    ("040700", "SERVIÇOS DE TERCEIROS NA PRODUÇÃO", "CUSTOS", "fornecedores"),
    ("050100", "PESSOAL ADMINISTRATIVO", "DESPESAS", "funcionarios"),
    ("050200", "DESPESAS COMERCIAIS (COMISSÕES, MARKETING)", "DESPESAS", ""),
    ("060100", "CONTABILIDADE", "DESPESAS", ""),
    ("060200", "SISTEMAS E TECNOLOGIA", "DESPESAS", ""),
    ("060400", "ADMINISTRATIVO", "DESPESAS", ""),
    ("060500", "OCUPAÇÃO (ALUGUEL, CONDOMÍNIO)", "DESPESAS", ""),
    ("060700", "SÓCIOS", "DESPESAS", ""),
    tuple(A_CLASSIFICAR),
    ("080100", "RECEITAS FINANCEIRAS", "FINANCEIRO", ""),
    ("080200", "DESPESAS FINANCEIRAS", "FINANCEIRO", ""),
    ("090100", "ENTRADAS NÃO OPERACIONAIS", "NAO_OPERACIONAL", ""),
    ("090200", "SAÍDAS NÃO OPERACIONAIS (INVESTIMENTOS, EMPRÉSTIMOS)", "NAO_OPERACIONAL", ""),
]
_INDUSTRIA_CONTAS = [
    ("Vendas", "010100"), ("Venda de Produtos", "010100"), ("ICMS", "020100"), ("IPI", "020100"),
    ("PIS", "020100"), ("COFINS", "020100"), ("Simples Nacional", "020100"), ("Devoluções", "020200"),
    ("Matéria-Prima", "040100"), ("Insumos", "040100"), ("Embalagens", "040100"), ("Salários Produção", "040200"),
    ("Salários", "040200"), ("FGTS", "040300"), ("INSS", "040300"), ("Vale Transporte", "040300"),
    ("Energia Elétrica", "040400"), ("Água", "040400"), ("Manutenção de Máquinas", "040500"),
    ("Frete", "040600"), ("Industrialização por Terceiros", "040700"), ("Salários Administrativo", "050100"),
    ("Comissões", "050200"), ("Marketing", "050200"), ("Contabilidade", "060100"), ("Sistemas", "060200"),
    ("Aluguel", "060500"), ("Pró-Labore", "060700"), ("Tarifas Bancárias", "080200"), ("Juros", "080200"),
    ("Compra de Máquinas", "090200"), ("Empréstimos", "090200"),
]

_NOTURNA_GRUPOS = [
    ("010100", "BILHETERIA (INGRESSOS E LISTAS)", "ENTRADAS", "recebimentos"),
    ("010200", "BAR, BEBIDAS E ALIMENTOS", "ENTRADAS", "recebimentos"),
    ("010300", "EVENTOS PRIVADOS E LOCAÇÃO DO ESPAÇO", "ENTRADAS", "recebimentos"),
    ("010400", "PATROCÍNIOS E PARCERIAS", "ENTRADAS", "recebimentos"),
    ("010500", "OUTRAS ENTRADAS OPERACIONAIS", "ENTRADAS", "recebimentos"),
    ("020100", "IMPOSTOS", "DEDUCOES", ""),
    ("020200", "TAXAS DE TICKETERIA E VENDAS ONLINE", "DEDUCOES", ""),
    ("040100", "ATRAÇÕES (CACHÊS DE ARTISTAS E DJS)", "CUSTOS", "fornecedores"),
    ("040200", "PRODUÇÃO (SOM, LUZ, PALCO E CENOGRAFIA)", "CUSTOS", "fornecedores"),
    ("040300", "EQUIPE DA NOITE (SEGURANÇA, BAR, PORTARIA, LIMPEZA)", "CUSTOS", "funcionarios"),
    ("040400", "BEBIDAS E INSUMOS DO BAR", "CUSTOS", "fornecedores"),
    ("040500", "ENCARGOS E BENEFÍCIOS", "CUSTOS", ""),
    ("040600", "DIREITOS AUTORAIS (ECAD)", "CUSTOS", ""),
    ("040700", "DIVULGAÇÃO DOS EVENTOS", "CUSTOS", "fornecedores"),
    ("050100", "PESSOAL ADMINISTRATIVO", "DESPESAS", "funcionarios"),
    ("060100", "CONTABILIDADE", "DESPESAS", ""),
    ("060200", "SISTEMAS E TECNOLOGIA", "DESPESAS", ""),
    ("060300", "MANUTENÇÃO DA CASA", "DESPESAS", "fornecedores"),
    ("060400", "ADMINISTRATIVO", "DESPESAS", ""),
    ("060500", "OCUPAÇÃO (ALUGUEL, CONDOMÍNIO, UTILIDADES)", "DESPESAS", ""),
    ("060600", "JURÍDICO, ALVARÁS E LICENÇAS", "DESPESAS", ""),
    ("060700", "SÓCIOS", "DESPESAS", ""),
    tuple(A_CLASSIFICAR),
    ("080100", "RECEITAS FINANCEIRAS", "FINANCEIRO", ""),
    ("080200", "DESPESAS FINANCEIRAS E TAXAS DE CARTÃO", "FINANCEIRO", ""),
    ("090100", "ENTRADAS NÃO OPERACIONAIS", "NAO_OPERACIONAL", ""),
    ("090200", "SAÍDAS NÃO OPERACIONAIS", "NAO_OPERACIONAL", ""),
]
_NOTURNA_CONTAS = [
    ("Ingressos", "010100"), ("Bilheteria", "010100"), ("Portaria", "010100"), ("Lista", "010100"),
    ("Sympla", "010100"), ("Ingresse", "010100"), ("Shotgun", "010100"), ("Camarote", "010100"),
    ("Bar", "010200"), ("Vendas Bar", "010200"), ("Consumação", "010200"), ("Cozinha", "010200"),
    ("Evento Privado", "010300"), ("Locação do Espaço", "010300"), ("Aniversário", "010300"),
    ("Patrocínio", "010400"), ("Parceria", "010400"),
    ("Simples Nacional", "020100"), ("ISS", "020100"), ("ICMS", "020100"),
    ("Taxa Ticketeria", "020200"), ("Taxa Sympla", "020200"), ("Taxa de Conveniência", "020200"),
    ("Cachê", "040100"), ("Cachê Artista", "040100"), ("DJ", "040100"), ("Atração", "040100"),
    ("Show", "040100"), ("Banda", "040100"), ("Som", "040200"), ("Iluminação", "040200"),
    ("Luz", "040200"), ("Palco", "040200"), ("Cenografia", "040200"), ("Produção", "040200"),
    ("Técnico de Som", "040200"), ("Segurança", "040300"), ("Bartender", "040300"), ("Garçons", "040300"),
    ("Hostess", "040300"), ("Freelancers", "040300"), ("Diaristas", "040300"), ("Limpeza", "040300"),
    ("Salários", "040300"), ("Gelo", "040400"), ("Bebidas", "040400"), ("Cerveja", "040400"),
    ("Destilados", "040400"), ("Insumos Bar", "040400"), ("Descartáveis", "040400"),
    ("Fornecedor de Bebidas", "040400"), ("FGTS", "040500"), ("INSS", "040500"), ("Vale Transporte", "040500"),
    ("Vale Refeição", "040500"), ("ECAD", "040600"), ("Direitos Autorais", "040600"), ("Divulgação", "040700"),
    ("Marketing", "040700"), ("Tráfego Pago", "040700"), ("Promoters", "040700"), ("Fotógrafo", "040700"),
    ("Influenciadores", "040700"), ("Contabilidade", "060100"), ("Sistemas", "060200"),
    ("Manutenção", "060300"), ("Telefonia", "060400"), ("Aluguel", "060500"), ("Energia Elétrica", "060500"),
    ("Água", "060500"), ("Alvará", "060600"), ("Licenças", "060600"), ("Advogados", "060600"),
    ("Bombeiros", "060600"), ("Pró-Labore", "060700"), ("Taxa de Cartão", "080200"),
    ("Tarifas Bancárias", "080200"), ("Antecipação de Recebíveis", "080200"),
]

_ARTISTAS_GRUPOS = [
    ("010100", "CACHÊS DE SHOWS E APRESENTAÇÕES", "ENTRADAS", "recebimentos"),
    ("010200", "DIREITOS AUTORAIS, STREAMING E ROYALTIES", "ENTRADAS", "recebimentos"),
    ("010300", "PUBLICIDADE, PATROCÍNIOS E PUBLIS", "ENTRADAS", "recebimentos"),
    ("010400", "PRODUTOS E MERCHANDISING", "ENTRADAS", "recebimentos"),
    ("010500", "OUTRAS ENTRADAS OPERACIONAIS", "ENTRADAS", "recebimentos"),
    ("020100", "IMPOSTOS", "DEDUCOES", ""),
    ("040100", "REPASSES A ARTISTAS", "CUSTOS", "fornecedores"),
    ("040200", "BANDA E MÚSICOS", "CUSTOS", "fornecedores"),
    ("040300", "EQUIPE TÉCNICA E PRODUÇÃO DE SHOWS", "CUSTOS", "fornecedores"),
    ("040400", "LOGÍSTICA DE TURNÊ (PASSAGENS, HOSPEDAGEM, TRANSPORTE)", "CUSTOS", "fornecedores"),
    ("040500", "LOCAÇÃO DE EQUIPAMENTOS", "CUSTOS", "fornecedores"),
    ("040600", "PRODUÇÃO MUSICAL E AUDIOVISUAL", "CUSTOS", "fornecedores"),
    ("040700", "MARKETING E DIVULGAÇÃO", "CUSTOS", "fornecedores"),
    ("040800", "COMISSÕES E AGENCIAMENTO", "CUSTOS", ""),
    ("050100", "PESSOAL", "DESPESAS", "funcionarios"),
    ("050200", "ENCARGOS E BENEFÍCIOS", "DESPESAS", ""),
    ("060100", "CONTABILIDADE", "DESPESAS", ""),
    ("060200", "SISTEMAS E TECNOLOGIA", "DESPESAS", ""),
    ("060400", "ADMINISTRATIVO", "DESPESAS", ""),
    ("060500", "OCUPAÇÃO (ESCRITÓRIO E ESTÚDIO)", "DESPESAS", ""),
    ("060600", "JURÍDICO E REGISTROS (OBRAS, MARCAS)", "DESPESAS", ""),
    ("060700", "SÓCIOS", "DESPESAS", ""),
    tuple(A_CLASSIFICAR),
    ("080100", "RECEITAS FINANCEIRAS", "FINANCEIRO", ""),
    ("080200", "DESPESAS FINANCEIRAS", "FINANCEIRO", ""),
    ("090100", "ENTRADAS NÃO OPERACIONAIS", "NAO_OPERACIONAL", ""),
    ("090200", "SAÍDAS NÃO OPERACIONAIS", "NAO_OPERACIONAL", ""),
]
_ARTISTAS_CONTAS = [
    ("Show", "010100"), ("Shows", "010100"), ("Cachê", "010100"), ("Apresentação", "010100"),
    ("Contratante", "010100"), ("Sinal de Show", "010100"), ("Direitos Autorais", "010200"),
    ("Streaming", "010200"), ("Royalties", "010200"), ("Spotify", "010200"), ("ECAD", "010200"),
    ("Distribuidora", "010200"), ("Publicidade", "010300"), ("Publi", "010300"), ("Patrocínio", "010300"),
    ("Merchandising", "010400"), ("Merch", "010400"), ("Produtos", "010400"),
    ("Simples Nacional", "020100"), ("ISS", "020100"), ("IRRF", "020100"),
    ("Repasse Artista", "040100"), ("Repasse", "040100"), ("Músicos", "040200"), ("Banda", "040200"),
    ("Backing Vocal", "040200"), ("Roadie", "040300"), ("Técnico de Som", "040300"), ("Produção", "040300"),
    ("Produtor", "040300"), ("Passagens", "040400"), ("Passagem Aérea", "040400"), ("Hospedagem", "040400"),
    ("Hotel", "040400"), ("Transporte", "040400"), ("Van", "040400"), ("Alimentação", "040400"),
    ("Diárias", "040400"), ("Combustível", "040400"), ("Locação de Equipamentos", "040500"),
    ("Backline", "040500"), ("Estúdio", "040600"), ("Gravação", "040600"), ("Mixagem", "040600"),
    ("Videoclipe", "040600"), ("Fotografia", "040600"), ("Marketing", "040700"), ("Divulgação", "040700"),
    ("Tráfego Pago", "040700"), ("Assessoria de Imprensa", "040700"), ("Comissão", "040800"),
    ("Agenciamento", "040800"), ("Empresário", "040800"), ("Salários", "050100"), ("FGTS", "050200"),
    ("INSS", "050200"), ("Contabilidade", "060100"), ("Sistemas", "060200"), ("Telefonia", "060400"),
    ("Aluguel", "060500"), ("Advogados", "060600"), ("Registro de Obras", "060600"), ("Pró-Labore", "060700"),
    ("Tarifas Bancárias", "080200"),
]

MODELOS_PRONTOS = {
    "Facilities (padrão Savian / HAGILI)": {"grupos": [list(g) for g in GRUPOS_PADRAO], "contas": _SERVICOS_CONTAS},
    "Casa noturna / Balada (ex.: Madame Satã)": {"grupos": [list(g) for g in _NOTURNA_GRUPOS],
                                                 "contas": _NOTURNA_CONTAS},
    "Promoção e agenciamento de artistas": {"grupos": [list(g) for g in _ARTISTAS_GRUPOS], "contas": _ARTISTAS_CONTAS},
    "Comércio / Varejo": {"grupos": [list(g) for g in _COMERCIO_GRUPOS], "contas": _COMERCIO_CONTAS},
    "Eventos e Festas": {"grupos": [list(g) for g in _EVENTOS_GRUPOS], "contas": _EVENTOS_CONTAS},
    "Indústria": {"grupos": [list(g) for g in _INDUSTRIA_GRUPOS], "contas": _INDUSTRIA_CONTAS},
}
MODELO_PADRAO = "Facilities (padrão Savian / HAGILI)"
NOMES_ANTIGOS = {"Serviços (padrão Savian) (pronto)": MODELO_PADRAO + " (pronto)"}


def completar_grupos(grupos):
    """Todo modelo precisa do grupo de transferências (fora da DRE) e do 'a classificar'."""
    gs = [list(g) for g in grupos]
    if not any(g[0] == A_CLASSIFICAR[0] for g in gs):
        gs.append(list(A_CLASSIFICAR))
    if not any(g[0] == IGNORAR for g in gs):
        gs.append(list(TRANSFERENCIAS))
    ordem = {s: i for i, (s, _) in enumerate(SECOES)}
    gs.sort(key=lambda g: (g[0] == IGNORAR, ordem.get(g[2], 99), g[0]))
    return gs


# ---------------------------------------------------------------- leitura do Excel
_RX_CODIGO = re.compile(r"^\s*(\d{6})\s*[-–—]\s*(.+?)\s*$")
_SECAO_POR_PREFIXO = {"01": "ENTRADAS", "02": "DEDUCOES", "03": "DEDUCOES", "04": "CUSTOS", "05": "DESPESAS",
                      "06": "DESPESAS", "07": "DESPESAS", "08": "FINANCEIRO", "09": "NAO_OPERACIONAL"}
_RELATORIOS = {"recebimentos", "fornecedores", "funcionarios", "transferencias"}


def _secao_do_titulo(txt):
    t = norm(txt)
    if t.startswith("(=)"):
        return None
    if "nao operacional" in t:
        return "NAO_OPERACIONAL"
    if "financeir" in t:
        return "FINANCEIRO"
    if "dedu" in t or "impostos" in t:
        return "DEDUCOES"
    if "custo" in t:
        return "CUSTOS"
    if "despesa" in t:
        return "DESPESAS"
    if "entrada" in t or "receita" in t:
        return "ENTRADAS"
    return None


def relatorio_sugerido(nome, secao):
    t = norm(nome)
    if secao == "ENTRADAS":
        return "recebimentos"
    if any(x in t for x in ("mao de obra", "pessoal", "equipe", "folha", "salario")):
        return "funcionarios"
    if secao == "CUSTOS" and any(x in t for x in ("fornecedor", "prestador", "material", "mercadoria", "materia",
                                                 "equipamento", "insumo", "terceiro", "frete", "locacao")):
        return "fornecedores"
    return ""


def ler_excel_modelo(arquivo):
    """Lê um Excel de DRE (um report de fechamento ou um modelo exportado pelo sistema).

    Procura linhas no formato '040100-NOME DO GRUPO' (nível 2, código terminando em 00) e
    '040101-Nome da conta' (nível 3), além dos títulos de seção como '(-) CUSTOS OPERACIONAIS'.
    Usa a primeira aba que tiver essas linhas. Devolve (grupos, contas).
    """
    wb = openpyxl.load_workbook(arquivo, data_only=True, read_only=True)
    for ws in wb.worksheets:
        grupos, contas, secao = {}, [], None
        for linha in ws.iter_rows(values_only=True):
            textos = [str(v).strip() for v in linha if isinstance(v, str) and v.strip()]
            rel = next((norm(t) for t in textos if norm(t) in _RELATORIOS), "")
            for t in textos:
                m = _RX_CODIGO.match(t)
                if m:
                    cod, nome = m.group(1), m.group(2)
                    if cod.endswith("00"):
                        sec = secao or _SECAO_POR_PREFIXO.get(cod[:2], "DESPESAS")
                        grupos[cod] = [cod, nome.upper(), sec, rel or relatorio_sugerido(nome, sec)]
                    else:
                        contas.append([nome, cod[:4] + "00"])
                    break
                s = _secao_do_titulo(t) if t.startswith("(") else None
                if s:
                    secao = s
                    break
        if grupos:
            # conta de um grupo que não apareceu: cria o grupo com o nome da própria conta
            for nome, cod in contas:
                if cod not in grupos:
                    sec = _SECAO_POR_PREFIXO.get(cod[:2], "DESPESAS")
                    grupos[cod] = [cod, nome.upper(), sec, relatorio_sugerido(nome, sec)]
            return completar_grupos(grupos.values()), contas
    raise ValueError("Não encontrei linhas no formato '040100-NOME DO GRUPO' em nenhuma aba do Excel. "
                     "Use um report de fechamento da Savian ou baixe o modelo em Excel aqui no sistema, "
                     "edite e importe de novo.")


# ---------------------------------------------------------------- exportação para edição
def exportar_excel_modelo(caminho_ou_buffer, nome, grupos, contas):
    """Gera um Excel no mesmo layout da DRE para o usuário editar e importar de volta."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MODELO DRE"
    azul = PatternFill("solid", fgColor="213A6B")
    ws.append([f"Modelo de DRE: {nome}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Edite os nomes, inclua ou exclua linhas mantendo o formato 'CÓDIGO-NOME'. Grupos (nível 2) "
               "terminam em 00; contas (nível 3) usam os mesmos 4 primeiros dígitos do grupo. A coluna Relatório "
               "aceita: recebimentos, fornecedores, funcionarios ou vazio."])
    ws.append([])
    ws.append(["Nível", "Descrição da Conta", "Relatório"])
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = azul
    por_grupo = {}
    for n, g in contas:
        por_grupo.setdefault(g, []).append(n)
    for sec, rotulo in SECOES:
        ws.append([1, rotulo, ""])
        ws.cell(ws.max_row, 2).font = Font(bold=True)
        for cod, gnome, gsec, rel in grupos:
            if gsec != sec:
                continue
            ws.append([2, f"{cod}-{gnome}", rel])
            ws.cell(ws.max_row, 2).font = Font(bold=True, color="213A6B")
            for i, n in enumerate(sorted(set(por_grupo.get(cod, []))), 1):
                ws.append([3, f"{cod[:4]}{i:02d}-{n}", ""])
    ws.column_dimensions["A"].width = 7
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 16
    wb.save(caminho_ou_buffer)


# ---------------------------------------------------------------- aplicação na empresa
def sugerir_pelo_modelo(categoria, cfg):
    """Grupo e nome de conta do modelo da empresa para uma categoria do OMIE (ou None).

    Primeiro o nome exato; depois a conta do modelo cujo nome aparece como palavra(s) na categoria
    (ex.: "Cachê DJ Fulano" -> grupo da conta "Cachê"), mantendo o nome da categoria como conta."""
    t = norm(categoria)
    contas = cfg.get("contas_modelo") or []
    for nome, grupo in contas:
        if norm(nome) == t:
            return grupo, nome
    codigos = {g[0] for g in cfg.get("grupos") or []}
    achados = [(len(norm(nome)), grupo) for nome, grupo in contas
               if len(norm(nome)) >= 2 and re.search(rf"\b{re.escape(norm(nome))}\b", t)
               and (not codigos or grupo in codigos)]
    if achados:
        return max(achados)[1], str(categoria).strip()
    return None
