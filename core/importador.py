"""Importa a planilha exportada do OMIE (Finanças > Movimentação da Conta Corrente)
como alternativa à API, no mesmo formato normalizado de omie.normalizar_mov_api."""
from datetime import date, datetime

import openpyxl

from .omie import _num


def _d(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    try:
        return datetime.strptime(str(v)[:10], "%d/%m/%Y").date()
    except ValueError:
        return None


def ler_exportacao_omie(arquivo):
    """arquivo: caminho ou file-like. Retorna (lancamentos, saldos)."""
    wb = openpyxl.load_workbook(arquivo, data_only=True, read_only=True)
    ws = None
    for w in wb.worksheets:
        for row in w.iter_rows(min_row=1, max_row=10, values_only=True):
            if row and "Categoria" in [str(x).strip() for x in row if x] and "Valor (R$)" in [str(x).strip() for x in row if x]:
                ws = w
                break
        if ws:
            break
    if ws is None:
        raise ValueError("Não encontrei a aba da movimentação (cabeçalho com 'Categoria' e 'Valor (R$)').")

    linhas = list(ws.iter_rows(values_only=True))
    i_hdr = next(i for i, r in enumerate(linhas) if r and "Valor (R$)" in [str(x).strip() for x in r if x])
    hdr = [str(x).strip() if x else "" for x in linhas[i_hdr]]
    ix = {h: i for i, h in enumerate(hdr) if h}

    def g(r, nome, padrao=""):
        i = ix.get(nome)
        return r[i] if i is not None and i < len(r) and r[i] is not None else padrao

    lancs, saldo_ini, saldo_fim, ultimo_saldo = [], None, None, {}
    for r in linhas[i_hdr + 1:]:
        if not r or not any(r):
            continue
        fav = str(g(r, "Cliente ou Fornecedor"))
        cc = g(r, "Conta Corrente") or "Conta Corrente"
        if fav.upper().startswith("SALDO"):
            if fav.upper() == "SALDO ANTERIOR" and saldo_ini is None:
                saldo_ini = _num(g(r, "Saldo (R$)", 0))
            else:
                saldo_fim = _num(g(r, "Saldo (R$)", 0))
            continue
        ultimo_saldo[cc] = _num(g(r, "Saldo (R$)", 0))
        origem = str(g(r, "Origem"))
        cat = str(g(r, "Categoria"))
        lancs.append({
            "data": _d(g(r, "Data")),
            "conta_corrente": cc,
            "favorecido": fav,
            "razao_social": g(r, "Cliente ou Fornecedor (Razão Social)"),
            "documento_favorecido": g(r, "Cliente ou Fornecedor (CNPJ/CPF)"),
            "categoria": cat,
            "cod_categoria": "",
            "transferencia": "transfer" in origem.lower() or "transfer" in cat.lower(),
            "valor": _num(g(r, "Valor (R$)", 0)),
            "projeto": g(r, "Projeto") or "",
            "situacao": g(r, "Situação"),
            "origem": origem,
            "tipo_documento": g(r, "Tipo de Documento"),
            "documento": g(r, "Documento") or g(r, "Nota Fiscal"),
            "observacoes": g(r, "Observações"),
            "id_omie": "",
        })
    contas = sorted({l["conta_corrente"] for l in lancs})
    saldos = []
    if len(contas) == 1:
        saldos.append({"conta": contas[0], "saldo_inicial": saldo_ini or 0.0,
                       "saldo_final": saldo_fim if saldo_fim is not None else ultimo_saldo.get(contas[0], 0.0)})
    return lancs, saldos
