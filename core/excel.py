"""Gera o Excel do fechamento no formato do modelo Savian: aba de REPORT com
fórmulas SUMIFS sobre a BASE CLASSIFICADA, mais SALDOS, RECEBIMENTOS etc.

A formatação reproduz a do modelo FECHAMENTO HAGILI: fonte Aptos Narrow, cabeçalhos de coluna
em negrito com borda fina (só o TOTAL em azul-marinho), seções e resultados com borda média em
cima e embaixo, grupos em cinza, contas recuadas, AV% em itálico e valores sem centavos."""
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, PatternFill, Side
from openpyxl.utils import get_column_letter as L

from . import dre as D
from .fluxo import (AZUL_DIGITADO, CINZA_GRUPO, FMT, FMT_AV, FMT_CENT, FMT_MES, GRADE, MARINHO, SECAO,
                    _cab, _fonte)
from .plano import IGNORAR

AMARELO_RATEIO = "FFF2CC"
_fino = Side(style="thin")
_medio = Side(style="medium")
TOTAL_MEDIO = Border(top=_medio, bottom=_medio, left=_medio, right=_medio)


def _fill(cor):
    return PatternFill("solid", fgColor=cor)


def gerar_excel(caminho, empresa, ano, mes, base, dre, saldos, cfg):
    wb = Workbook()
    ws = wb.active
    modo = {"projeto": "POR PROJETO", "conta": "POR CONTA"}.get(cfg.get("modo_colunas"))
    ws.title = f"REPORT {mes:02d} {ano}" + (f" - {modo}" if modo else "")
    n_base = len(base) + 1
    _aba_base(wb, base)
    _aba_report(ws, empresa, ano, mes, dre, n_base, modo)
    _aba_saldos(wb, saldos, ano, mes)
    rotulo_rec = "PROJETO" if cfg.get("modo_colunas") == "projeto" else None
    _aba_quadro(wb, "RECEBIMENTOS", D.recebimentos(base, cfg), "RECEBIDO", rotulo_rec, com_av=True)
    _aba_quadro(wb, "FORNECEDORES", D.pagamentos(base, cfg, "fornecedores"), "VALOR PAGO", "FORNECEDOR")
    _aba_quadro(wb, "FUNCIONARIOS", D.pagamentos(base, cfg, "funcionarios"), "VALOR PAGO", "FUNCIONÁRIO")
    _aba_transferencias(wb, D.transferencias(base))
    _aba_indicadores(wb, D.indicadores(dre), ano, mes)
    wb.move_sheet("BASE CLASSIFICADA", offset=len(wb.sheetnames))
    wb.save(caminho)
    return caminho


def _aba_report(ws, empresa, ano, mes, dre, n_base, modo):
    cols, agrup = dre["colunas"], dre["agrupamentos"]
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 54
    ws.row_dimensions[2].height = 45
    titulo = "CashFlowReport" + (f" - {modo.title()}" if modo else "")
    _cab(ws, 2, 2, titulo, fill=None, cor=None, borda=None).alignment = Alignment(horizontal="center",
                                                                                   vertical="center")
    c = ws.cell(3, 2, "Descrição da Conta")
    c.font = _fonte(bold=True)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "C4"

    # layout das colunas: [coluna, coluna, ..., subtotal agrupamento], ..., TOTAL
    blocos, usados = [], set()
    for g, membros in agrup.items():
        for m in membros:
            blocos.append(("col", m))
            usados.add(m)
        blocos.append(("sub", g))
    for c in cols:
        if c not in usados:
            blocos.append(("col", c))
    blocos.append(("tot", f"TOTAL {empresa['nome']}".upper()))

    pos, cpos = {}, 3
    competencia = date(ano, mes, 1)
    for tipo, nome in blocos:
        pos[(tipo, nome)] = cpos
        ws.merge_cells(start_row=2, start_column=cpos, end_row=2, end_column=cpos + 1)
        destaque = tipo != "col"
        _cab(ws, 2, cpos, nome if tipo != "sub" else f"TOTAL {nome}".upper(),
             fill=MARINHO if destaque else None, cor="FFFFFF" if destaque else None)
        ws.cell(2, cpos + 1).border = Border(top=_fino, bottom=_fino, right=_fino)
        for col, v, fmt in ((cpos, competencia, FMT_MES), (cpos + 1, "AV", "General")):
            h = ws.cell(3, col, v)
            h.font = _fonte(bold=True)
            h.alignment = Alignment(horizontal="center")
            h.number_format = fmt
        ws.column_dimensions[L(cpos)].width = 13
        ws.column_dimensions[L(cpos + 1)].width = 6.57
        ws.column_dimensions[L(cpos + 2)].width = 1.14
        cpos += 3
    col_total = pos[("tot", blocos[-1][1])]
    col_conf = cpos
    _cab(ws, 2, col_conf, "CONFERÊNCIA (TOTAL − BASE)", fill=None, cor=None, borda=None).font = _fonte(bold=True,
                                                                                                         size=9)
    ws.column_dimensions[L(col_conf)].width = 14
    # colunas que recebem borda de seção / cinza de grupo: B até o AV do TOTAL (inclui os espaçadores)
    faixa = range(2, col_total + 2)

    rng_v = f"'BASE CLASSIFICADA'!$K$2:$K${n_base}"
    rng_c = f"'BASE CLASSIFICADA'!$G$2:$G${n_base}"
    rng_col = f"'BASE CLASSIFICADA'!$I$2:$I${n_base}"

    linhas = dre["linhas"]
    r0 = 4
    row_of = {i: r0 + i for i in range(len(linhas))}
    secao_row, grupos_por_secao, contas_por_grupo = {}, {}, {}
    for i, l in enumerate(linhas):
        if l["tipo"] == "secao":
            secao_row[l["secao"]] = row_of[i]
            grupos_por_secao[l["secao"]] = []
        elif l["tipo"] == "grupo":
            grupos_por_secao[l["secao"]].append(row_of[i])
            contas_por_grupo[row_of[i]] = []
            g_atual = row_of[i]
        elif l["tipo"] == "conta":
            contas_por_grupo[g_atual].append(row_of[i])

    from .plano import RESULTADOS, SECOES
    for i, l in enumerate(linhas):
        r = row_of[i]
        ws.cell(r, 1, l["nivel"]).font = _fonte(size=8, color="FFFFFF")
        c = ws.cell(r, 2, l["descricao"])
        bold = l["tipo"] != "conta"
        c.font = _fonte(bold=bold)
        if l["tipo"] == "conta":
            c.alignment = Alignment(indent=1)
        for (tipo, nome), cp in pos.items():
            X = L(cp)
            if tipo == "col":
                if l["tipo"] == "conta":
                    f = f"=SUMIFS({rng_v},{rng_c},$B{r},{rng_col},\"{nome}\")"
                elif l["tipo"] == "grupo":
                    rs = contas_por_grupo[r]
                    f = f"=SUM({X}{rs[0]}:{X}{rs[-1]})" if rs else 0
                elif l["tipo"] == "secao":
                    rs = grupos_por_secao.get(l["secao"], [])
                    f = "=" + "+".join(f"{X}{x}" for x in rs) if rs else 0
                else:
                    if l["secao"] in RESULTADOS:
                        secs = RESULTADOS[l["secao"]][1]
                    else:
                        secs = [s for s, _ in SECOES]
                    f = "=" + "+".join(f"{X}{secao_row[s]}" for s in secs if s in secao_row)
            elif tipo == "sub":
                membros = agrup[nome]
                f = "=" + "+".join(f"{L(pos[('col', m)])}{r}" for m in membros)
            else:
                f = "=" + "+".join(f"{L(pos[('col', m)])}{r}" for m in cols) if cols else 0
            cv = ws.cell(r, cp, f)
            cv.number_format = FMT
            cv.font = _fonte(bold=bold)
            ent = secao_row.get("ENTRADAS", r0)
            ca = ws.cell(r, cp + 1, f"=IFERROR({X}{r}/{X}${ent},0)")
            ca.number_format = FMT_AV
            ca.font = _fonte(bold=bold, size=10, italic=True)
        if l["tipo"] in ("secao", "resultado"):
            for cc in faixa:
                ws.cell(r, cc).border = SECAO
        elif l["tipo"] == "grupo":
            for cc in faixa:
                ws.cell(r, cc).fill = _fill(CINZA_GRUPO)
        if l["tipo"] == "conta":
            T = L(col_total)
            cc = ws.cell(r, col_conf, f"=ROUND({T}{r}-SUMIFS({rng_v},{rng_c},$B{r}),2)")
            cc.number_format = FMT_CENT
            cc.font = _fonte(size=10, italic=True, color="808080")

    fim = r0 + len(linhas)
    T = L(col_total)
    liq_row = r0 + next(i for i, l in enumerate(linhas) if l["descricao"] == D.RESULTADO_LIQUIDO)
    rng_g = f"'BASE CLASSIFICADA'!$F$2:$F${n_base}"
    conf = [
        ("Conferência: TOTAL − soma da BASE CLASSIFICADA (deve ser 0)",
         f"=ROUND({T}{liq_row}-SUMIFS({rng_v},{rng_g},\"<>{IGNORAR}\"),2)"),
        ("Conferência: resultado líquido + transferências − variação de caixa (deve ser 0)",
         f"=ROUND({T}{liq_row}+SUMIFS({rng_v},{rng_g},\"{IGNORAR}\")-SALDOS!E{{TOT}},2)"),
    ]
    for k, (txt, f) in enumerate(conf):
        ws.cell(fim + 1 + k, 2, txt).font = _fonte(size=10, italic=True)
        x = ws.cell(fim + 1 + k, col_total, f)
        x.number_format = FMT_CENT
        x.font = _fonte(size=10, italic=True)
    ws._savian_conf_cell = ws.cell(fim + 2, col_total)
    ws.cell(fim + 4, 2, "Notas:").font = _fonte(bold=True, size=10)
    notas = [
        "• Valores puxados por SUMIFS da aba BASE CLASSIFICADA (cada lançamento do extrato OMIE + linhas de rateio).",
        "• Categorias com % no nome (ex.: \"Combustível (98%); Tarifas (2%)\") foram divididas entre as contas pelos percentuais.",
        "• DRE financeira (regime de caixa): considera apenas lançamentos realizados; transferências entre contas ficam fora.",
    ]
    for i, n in enumerate(notas):
        ws.cell(fim + 5 + i, 2, n).font = _fonte(size=10)


def _aba_base(wb, base):
    ws = wb.create_sheet("BASE CLASSIFICADA")
    hdr = ["Data", "Cliente ou Fornecedor", "Conta Corrente", "Categoria (original)", "Componente",
           "Grupo", "Conta do Fechamento", "Projeto (original)", "Coluna (fechamento)", "Situação", "Valor (R$)",
           "Tipo", "Observações"]
    larg = [11, 38, 18, 34, 26, 9, 40, 32, 26, 13, 13, 11, 60]
    for i, (h, w) in enumerate(zip(hdr, larg), start=1):
        _cab(ws, 1, i, h, borda=None)
        ws.column_dimensions[L(i)].width = w
    for r, b in enumerate(base, start=2):
        vals = [b["data"], b["favorecido"], b["conta_corrente"], b["categoria"], b["componente"], b["grupo"],
                b["conta_fechamento"], b["projeto"], b["coluna"], b["situacao"], b["valor"], b["tipo"],
                (b["observacoes"] or "")[:500]]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(r, c, v)
            cell.font = _fonte()
            if c == 1:
                cell.number_format = "mm-dd-yy"
            if c == 11:
                cell.number_format = FMT_CENT
            if b["tipo"] != "Lançamento":
                cell.fill = _fill(AMARELO_RATEIO)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{L(len(hdr))}{len(base) + 1}"


def _aba_saldos(wb, saldos, ano, mes):
    import calendar
    ws = wb.create_sheet("SALDOS")
    ult = calendar.monthrange(ano, mes)[1]
    ini = f"{calendar.monthrange(ano - (mes == 1), 12 if mes == 1 else mes - 1)[1]:02d}/{12 if mes == 1 else mes - 1:02d}/{ano - (mes == 1)}"
    ws.row_dimensions[2].height = 30
    for i, h in enumerate(["CONTA", f"SALDO INICIAL {ini}", f"SALDO FINAL {ult:02d}/{mes:02d}/{ano}", "VARIAÇÃO"], start=2):
        _cab(ws, 2, i, h, borda=None).alignment = Alignment(horizontal="center", wrap_text=True)
    for r, s in enumerate(saldos, start=3):
        ws.cell(r, 2, s["conta"]).font = _fonte()
        for c, v in ((3, s["saldo_inicial"]), (4, s["saldo_final"]), (5, f"=D{r}-C{r}")):
            x = ws.cell(r, c, v)
            x.number_format = FMT_CENT
            x.font = _fonte(color=AZUL_DIGITADO if c < 5 else None)
    tot = 3 + len(saldos)
    ws.cell(tot, 2, "TOTAL").font = _fonte(bold=True)
    for c in (3, 4, 5):
        x = ws.cell(tot, c, f"=SUM({L(c)}3:{L(c)}{tot - 1})" if saldos else 0)
        x.number_format = FMT_CENT
        x.font = _fonte(bold=True)
    for col, w in zip("ABCDE", (2, 28, 20, 20, 20)):
        ws.column_dimensions[col].width = w
    rep = wb.worksheets[0]
    cel = getattr(rep, "_savian_conf_cell", None)
    if cel is not None:
        cel.value = cel.value.replace("{TOT}", str(tot))


def _titulo_quadro(ws, row, texto, ncols):
    """Linha de título do quadro (como "RECEBIMENTOS POR PROJETO" no modelo): mesclada, negrito, borda fina."""
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=1 + ncols)
    _cab(ws, row, 2, texto, fill=None, cor=None).alignment = Alignment(horizontal="center")
    for c in range(3, 2 + ncols):
        ws.cell(row, c).border = Border(top=_fino, bottom=_fino, right=_fino if c == 1 + ncols else None)


def _cab_quadro(ws, row, textos):
    for i, h in enumerate(textos, start=2):
        _cab(ws, row, i, h).alignment = Alignment(horizontal="center")


def _linha_total(ws, row, valores, fmts):
    """Linha de total do quadro: negrito com borda média, como "TOTAL RECEBIMENTO" no modelo."""
    for i, (v, fmt) in enumerate(zip(valores, fmts), start=2):
        x = ws.cell(row, i, v)
        x.font = _fonte(bold=True)
        x.border = TOTAL_MEDIO
        if fmt:
            x.number_format = fmt


def _rotulo_item(bloco, padrao):
    if padrao:
        return padrao
    if "POR CATEGORIA" in bloco:
        return "CATEGORIA"
    return "CLIENTE"


def _aba_quadro(wb, nome, quadro, rotulo_valor, rotulo_item=None, com_av=False):
    ws = wb.create_sheet(nome)
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 38
    ws.column_dimensions["C"].width = 13
    ws.column_dimensions["D"].width = 7.14
    ncols = 3 if com_av else 2
    r = 2
    total_rows = []
    for bloco, itens in quadro.items():
        _titulo_quadro(ws, r, bloco, ncols)
        _cab_quadro(ws, r + 1, [_rotulo_item(bloco, rotulo_item), rotulo_valor] + (["AV%"] if com_av else []))
        r += 2
        ini = r
        fim = r + len(itens)
        for item, v in itens:
            ws.cell(r, 2, item)
            ws.cell(r, 3, v).number_format = FMT_CENT
            if com_av:
                ws.cell(r, 4, f"=IFERROR(C{r}/$C${fim},0)").number_format = "0.0%"
            for c in range(2, 2 + ncols):
                ws.cell(r, c).font = _fonte()
                ws.cell(r, c).border = GRADE
            r += 1
        vals = [f"TOTAL {bloco}", f"=SUM(C{ini}:C{r - 1})"] + ([f"=SUM(D{ini}:D{r - 1})"] if com_av else [])
        _linha_total(ws, r, vals, [None, FMT_CENT, "0.0%"])
        total_rows.append(r)
        r += 2
    if len(total_rows) > 1:
        _linha_total(ws, r, ["TOTAL GERAL", "=" + "+".join(f"C{t}" for t in total_rows)], [None, FMT_CENT])
    if not quadro:
        ws.cell(2, 2, "Sem lançamentos no mês para este quadro.").font = _fonte()


def _aba_transferencias(wb, transf):
    ws = wb.create_sheet("TRANSFERENCIAS")
    _titulo_quadro(ws, 2, "TRANSFERÊNCIAS ENTRE CONTAS DA EMPRESA", 3)
    _cab_quadro(ws, 3, ["CONTA", "TRANSFERÊNCIA ENVIADA", "TRANSFERÊNCIA RECEBIDA"])
    r = 4
    for conta, v in transf.items():
        ws.cell(r, 2, conta)
        ws.cell(r, 3, round(v["enviada"], 2)).number_format = FMT_CENT
        ws.cell(r, 4, round(v["recebida"], 2)).number_format = FMT_CENT
        for c in (2, 3, 4):
            ws.cell(r, c).font = _fonte()
            ws.cell(r, c).border = GRADE
        r += 1
    somas = [f"=SUM({L(c)}4:{L(c)}{r - 1})" if transf else 0 for c in (3, 4)]
    _linha_total(ws, r, ["TOTAL TRANSFERIDO"] + somas, [None, FMT_CENT, FMT_CENT])
    for col, w in zip("ABCD", (2, 34, 20, 20)):
        ws.column_dimensions[col].width = w


def _aba_indicadores(wb, ind, ano, mes):
    ws = wb.create_sheet("INDICADORES")
    _cab_quadro(ws, 2, ["Indicador", date(ano, mes, 1), "% dos recebimentos"])
    ws.cell(2, 3).number_format = FMT_MES
    for r, (k, v) in enumerate(ind.items(), start=3):
        destaque = k.startswith("Resultado") or k in ("Entradas", "Receita Líquida")
        ws.cell(r, 2, k)
        ws.cell(r, 3, v["valor"]).number_format = FMT_CENT
        ws.cell(r, 4, v["pct"]).number_format = "0.0%"
        for c in (2, 3, 4):
            ws.cell(r, c).font = _fonte(bold=destaque)
            ws.cell(r, c).border = GRADE
    for col, w in zip("ABCD", (2, 30, 16, 16)):
        ws.column_dimensions[col].width = w
