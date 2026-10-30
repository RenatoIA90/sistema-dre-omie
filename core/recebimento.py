"""Relatório de recebimento (Finanças › Contas a Receber do OMIE).

Segue a planilha modelo RELATORIO DE RECEBIMENTO LETS: as 11 colunas que ela mantém da exportação do
OMIE e o mesmo filtro, que é:
  - títulos recebidos com previsão de recebimento dentro do período;
  - títulos em aberto (a vencer, vence hoje, atrasados e recebidos parcialmente) com previsão até o fim do
    período, inclusive os atrasados de meses anteriores;
  - cancelados ficam de fora.
Ordem: situação (A vencer, Vence hoje, Atrasado, Recebido, Recebido Parcialmente) e previsão da mais recente
para a mais antiga.

A base vem da API (financas/mf ListarMovimentos para situação e valores, ListarContasReceber para NF,
retenções e observação, cadastro de clientes para o nome fantasia) ou da planilha exportada do OMIE.
"""
import html
import os
from datetime import date

import openpyxl
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, GradientFill, Side
from openpyxl.utils import get_column_letter as L

from .omie import _data, _num, _pega

A_VENCER, VENCE_HOJE, ATRASADO = "A vencer", "Vence hoje", "Atrasado"
RECEBIDO, PARCIAL = "Recebido", "Recebido Parcialmente"
SITUACOES = [A_VENCER, VENCE_HOJE, ATRASADO, RECEBIDO, PARCIAL]
EM_ABERTO = {A_VENCER, VENCE_HOJE, ATRASADO, PARCIAL}

# (chave, cabeçalho, largura) na ordem da planilha modelo
COLUNAS = [
    ("situacao", "Situação", 13.28515625),
    ("nf", "NF", 8.7109375),
    ("cliente", "Cliente (Nome Fantasia)", 76.42578125),
    ("previsao", "Previsão de Recebimento", 38.5703125),
    ("ultimo_recebimento", "Último Recebimento", 31.140625),
    ("valor_conta", "Valor da Conta", 25.140625),
    ("valor_liquido", "Valor Líquido", 23.7109375),
    ("impostos", "Impostos Retidos", 28.28515625),
    ("valor_receber", "Valor a Receber", 26.5703125),
    ("categoria", "Categoria", 38.85546875),
    ("observacao", "Observação", 411.5703125),
]
DATAS = {"previsao", "ultimo_recebimento"}
VALORES = {"valor_conta", "valor_liquido", "impostos", "valor_receber"}
IMPOSTOS = ("pis", "cofins", "csll", "ir", "iss", "inss")

# formatação da planilha modelo
FMT_VALOR = '#,##0.00;[Red]\\(#,##0.00\\)'
FMT_DATA = "mm-dd-yy"  # formato de data curta do Excel (aparece dd/mm/aaaa no Excel em português)
AZUL_CAB = "005795"
CINZA_CAB = "F0F0F0"
_fino = Side(style="thin")
LOGO = os.path.join(os.path.dirname(os.path.dirname(__file__)), "assets", "logo_savian.png")


def situacao(status_omie, previsao, pago, aberto, hoje=None):
    """Situação como a tela do OMIE mostra; None para cancelado."""
    if "CANCEL" in str(status_omie or "").upper():
        return None
    if aberto <= 0.005 and (pago > 0 or str(status_omie).upper() == "RECEBIDO"):
        return RECEBIDO
    if pago > 0.005:
        return PARCIAL
    hoje = hoje or date.today()
    if previsao and previsao < hoje:
        return ATRASADO
    return VENCE_HOJE if previsao == hoje else A_VENCER


def _retencoes(cr, det):
    """Soma dos impostos retidos do título (só os marcados como retidos)."""
    tot = 0.0
    for imp in IMPOSTOS:
        if str(cr.get(f"retem_{imp}", "")).upper() == "S":
            tot += _num(cr.get(f"valor_{imp}", 0))
        elif str(det.get(f"cRet{imp.upper()}", "")).upper() == "S":
            tot += _num(det.get(f"nValor{imp.upper()}", 0))
    return round(tot, 2)


def baixar_titulos(cli, progresso=None):
    """Todos os títulos a receber da empresa, já normalizados nas colunas do relatório."""
    passos = [("Baixando contas a receber (situação e valores)...", cli.movimentos_receber),
              ("Baixando contas a receber (NF e observações)...", cli.contas_receber),
              ("Baixando cadastro de clientes...", cli.clientes),
              ("Baixando categorias...", cli.categorias)]
    res = []
    for i, (texto, func) in enumerate(passos):
        if progresso:
            progresso(i / len(passos), texto)
        res.append(func())
    movs, crs, clientes, cats = res
    if progresso:
        progresso(1.0, "Contas a receber baixadas")
    por_cr = {str(c.get("codigo_lancamento_omie")): c for c in crs}
    nomes = {str(c.get("codigo_cliente")): html.unescape(str(c.get("nome_fantasia") or c.get("razao_social") or ""))
             for c in clientes}
    nome_cat = {str(c.get("codigo")): html.unescape(str(c.get("descricao") or "")) for c in cats}

    titulos = []
    for m in movs:
        det, resumo = m.get("detalhes") or {}, m.get("resumo") or {}
        # o ListarMovimentos também devolve as baixas e os lançamentos avulsos da conta corrente;
        # o relatório usa só os títulos (os que têm resumo de liquidação)
        if "cLiquidado" not in resumo or str(det.get("cGrupo", "CONTA_A_RECEBER")) != "CONTA_A_RECEBER":
            continue
        cod = str(det.get("nCodTitulo"))
        cr = por_cr.get(cod, {})
        previsao = _data(_pega(det, "dDtPrevisao", "dDtVenc") or cr.get("data_previsao"))
        valor = _num(_pega(det, "nValorTitulo", padrao=0) or cr.get("valor_documento", 0))
        pago = _num(resumo.get("nValPago", 0))
        aberto = _num(resumo.get("nValAberto", 0))
        sit = situacao(det.get("cStatus") or cr.get("status_titulo"), previsao, pago, aberto)
        if sit is None:
            continue
        impostos = _retencoes(cr, det)
        liquido = round(valor - impostos, 2)
        if sit == RECEBIDO:
            a_receber = 0.0
        elif pago <= 0.005:
            a_receber = liquido
        else:
            a_receber = round(min(aberto, liquido - pago) if impostos else aberto, 2)
        cod_cat = str(det.get("cCodCateg") or cr.get("codigo_categoria") or "")
        cliente = nomes.get(str(det.get("nCodCliente") or cr.get("codigo_cliente_fornecedor")), "")
        titulos.append({
            "situacao": sit,
            "nf": str(_pega(det, "cNumDocFiscal", padrao="") or _pega(cr, "numero_documento_fiscal", padrao="")) or None,
            "cliente": cliente or _pega(det, "cCPFCNPJCliente", padrao=""),
            "previsao": previsao,
            "ultimo_recebimento": _data(det.get("dDtPagamento")) if pago > 0.005 else None,
            "valor_conta": valor,
            "valor_liquido": liquido,
            "impostos": impostos,
            "valor_receber": a_receber,
            "categoria": nome_cat.get(cod_cat, cod_cat),
            "observacao": _pega(cr, "observacao", padrao="") or None,
        })
    return titulos


def ler_exportacao(arquivo):
    """Planilha exportada de Finanças › Contas a Receber do OMIE (ou já no formato do modelo)."""
    wb = openpyxl.load_workbook(arquivo, data_only=True, read_only=True)
    for ws in wb.worksheets:
        linhas = list(ws.iter_rows(values_only=True))
        i_hdr = next((i for i, r in enumerate(linhas[:15])
                      if r and {"Situação", "Valor a Receber"} <= {str(x).strip() for x in r if x}), None)
        if i_hdr is not None:
            break
    else:
        raise ValueError("Não encontrei o cabeçalho das contas a receber (colunas 'Situação' e 'Valor a Receber').")
    ix = {str(h).strip(): i for i, h in enumerate(linhas[i_hdr]) if h}
    nomes = {"nf": ("NF", "Nota Fiscal/Cupom Fiscal", "Nota Fiscal")}
    nomes.update({k: (titulo,) for k, titulo, _ in COLUNAS if k != "nf"})

    def g(r, chave):
        for nome in nomes[chave]:
            i = ix.get(nome)
            if i is not None and i < len(r) and r[i] not in (None, ""):
                return r[i]
        return None

    titulos = []
    for r in linhas[i_hdr + 1:]:
        if not r or not g(r, "situacao"):
            continue
        sit = str(g(r, "situacao")).strip()
        if "cancel" in sit.lower():
            continue
        t = {k: g(r, k) for k, _, _ in COLUNAS}
        t["situacao"] = next((s for s in SITUACOES if s.lower() == sit.lower()), sit)
        for k in DATAS:
            t[k] = _data(t[k]) if t[k] else None
        for k in VALORES:
            t[k] = _num(t[k] or 0)
        t["nf"] = str(t["nf"]) if t["nf"] is not None else None
        titulos.append(t)
    return titulos


def filtrar(titulos, inicio, fim, situacoes=None):
    """Mesmo filtro e ordem da planilha modelo."""
    sel = []
    for t in titulos:
        if situacoes and t["situacao"] not in situacoes:
            continue
        p = t["previsao"]
        if p is None or p > fim:
            continue
        if t["situacao"] not in EM_ABERTO and p < inicio:
            continue
        sel.append(t)
    ordem = {s: i for i, s in enumerate(SITUACOES)}
    # empates (mesma situação e previsão) ficam na ordem em que o OMIE devolveu, como no modelo
    return sorted(sel, key=lambda t: (ordem.get(t["situacao"], len(ordem)), -t["previsao"].toordinal()))


def resumo(titulos):
    abertos = [t for t in titulos if t["situacao"] in EM_ABERTO]
    return {
        "titulos": len(titulos),
        "recebido": sum(t["valor_liquido"] for t in titulos if t["situacao"] == RECEBIDO),
        "a_receber": sum(t["valor_receber"] for t in abertos),
        "atrasado": sum(t["valor_receber"] for t in abertos if t["situacao"] in (ATRASADO, PARCIAL)),
        "a_vencer": sum(t["valor_receber"] for t in abertos if t["situacao"] in (A_VENCER, VENCE_HOJE)),
    }


def gerar_excel(caminho, empresa, inicio, fim, titulos):
    """Excel com a formatação da planilha modelo (aba financas, cabeçalho na linha 3, filtro e painel congelado)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "financas"
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 39.95
    ws.row_dimensions[2].height = 24.95
    ws.row_dimensions[3].height = 18.75
    ultima = L(len(COLUNAS))
    ws.merge_cells(f"B1:{ultima}1")
    ws.merge_cells(f"B2:{ultima}2")
    if os.path.exists(LOGO):
        img = Image(LOGO)
        img.width, img.height = 207, 64
        ws.add_image(img, "A1")
    # o modelo deixa as linhas 1 e 2 em branco ao lado do logo; aqui elas identificam o relatório
    c = ws["B1"]
    c.value = f"RELATÓRIO DE RECEBIMENTO — {empresa['nome']}"
    c.font = Font(name="Calibri", size=16, bold=True, color=AZUL_CAB)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=14)
    c = ws["B2"]
    c.value = (f"Recebidos com previsão de {inicio:%d/%m/%Y} a {fim:%d/%m/%Y} · em aberto com previsão até "
               f"{fim:%d/%m/%Y} · gerado em {date.today():%d/%m/%Y}")
    c.font = Font(name="Calibri", size=11, color="595959")
    c.alignment = Alignment(horizontal="left", vertical="center", indent=18)

    cab_fill = GradientFill(type="linear", degree=0, stop=(CINZA_CAB, CINZA_CAB))
    for j, (_, titulo, largura) in enumerate(COLUNAS, 1):
        ws.column_dimensions[L(j)].width = largura
        c = ws.cell(3, j, titulo)
        c.font = Font(name="Calibri", size=14, bold=True, color=AZUL_CAB)
        c.fill = cab_fill
        c.alignment = Alignment(horizontal="left", vertical="top")
        c.border = Border(top=_fino, bottom=_fino)

    fonte = Font(name="Calibri", size=11, color="000000")
    for i, t in enumerate(titulos, 4):
        for j, (k, _, _) in enumerate(COLUNAS, 1):
            c = ws.cell(i, j, t.get(k))
            c.font = fonte
            if k in DATAS:
                c.number_format = FMT_DATA
                c.alignment = Alignment(horizontal="center", vertical="top")
            elif k in VALORES:
                c.number_format = FMT_VALOR
                c.alignment = Alignment(horizontal="right", vertical="top")
            else:
                c.number_format = "0"
                c.alignment = Alignment(horizontal="left", vertical="top")

    fim_linha = max(3 + len(titulos), 4)
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{ultima}{fim_linha}"
    ws.print_title_rows = "1:3"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_margins.left = ws.page_margins.right = 0
    ws.oddFooter.left.text = "&BFinanças - Contas a Receber"
    ws.oddFooter.right.text = "Página &P de &N"
    wb.properties.title = wb.properties.subject = "Finanças - Contas a Receber"
    wb.save(caminho)
    return caminho
